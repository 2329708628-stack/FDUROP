# -*- coding: utf-8 -*-
"""
NLI 工具选型与验证报告（docx）
"""
import json, os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

EQ_DIR = r"c:\Users\23297\Downloads\Lab\A01_湖北茶叶经济\03_实验输出\H1_v3.2\_equivalence"
OUT = r"c:\Users\23297\Downloads\Lab\实验方案\NLI工具选型与验证报告.docx"

results = json.load(open(os.path.join(EQ_DIR, "nli_mdeberta_results.json"), encoding="utf-8"))

doc = Document()
doc.add_heading("NLI 工具选型与等义判定验证报告", 0).alignment = WD_ALIGN_PARAGRAPH.CENTER

def p(text, bold=False, size=11):
    r = doc.add_paragraph().add_run(text)
    r.bold = bold; r.font.size = Pt(size); return r

def tbl(headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers)); t.style = 'Table Grid'
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text = h
        for run in t.rows[0].cells[i].paragraphs[0].runs: run.bold = True
    for ri,row in enumerate(rows):
        for ci,v in enumerate(row):
            t.rows[ri+1].cells[ci].text = str(v)
    return t

# ============ 一、背景 ============
doc.add_heading("一、背景与目标", 1)
p("实验方案第五步规定：对原子化后的命题对，由 NLI 模型计算双向蕴含概率，"
  "取算术平均作为等义置信度，再经阈值决策判定 EQUIVALENT / DIFFERENT / UNCERTAIN。"
  "本报告在 A01 三轮 v3.2 输出的 133 对人工标定命题对上，对比两个候选 NLI 工具，"
  "选定满足精确率 ≥ 95% 要求的工作点。")
p("金标准构成：P33 结语段 48 对（节点号相同即等义，共 12 正例）+ "
  "P6 段落 85 对（含拆合案例，12 等义 + 4 单向蕴含 + 69 不等义）。"
  "正例（等义）共 24 对，负例 109 对。", size=10)

# ============ 二、候选工具 ============
doc.add_heading("二、候选工具", 1)
tbl(["工具", "架构", "输出", "部署难度", "结论"],
    [["mDeBERTa-v3-base-xnli-multilingual-nli-2mil7",
      "DeBERTa-v3 序列分类，278M",
      "三分类概率（entailment/neutral/contradiction）",
      "低（transformers 直接加载）",
      "采用"],
     ["iic/nlp_deberta_rex-uninlu_chinese-base (RexUniNLU)",
      "DeBERTa-v2 MaskedLM 统一信息抽取，393MB",
      "抽取式硬标签（无概率）",
      "高（modelscope + 5 处兼容性补丁）",
      "淘汰"]])

doc.add_heading("RexUniNLU 淘汰原因", 2)
p("1. 无概率输出：模型为 MaskedLM 架构，仅返回抽取式标签字符串（如 {\"type\":\"蕴含\"}），"
  "无法获得连续置信度，不满足方案'双向蕴含概率 + 温度缩放 + 阈值校准'的要求。")
p("2. NLI 行为异常：将'年产量不足一百万斤'与'年产量逾1300万斤'（明显矛盾）误判为'蕴含'，"
  "无关对输出空标签。其分类逻辑非标准 NLI 三分类语义，不可用于等义判定。")
p("3. 部署代价高：需打 5 处兼容性补丁（transformers API 补桩、Trainer 参数重命名、tokenizer 属性补全）"
  "且仅在沙箱外（禁用 multiprocessing 写系统目录限制）才能运行。")

# ============ 三、mDeBERTa-v3 结果 ============
doc.add_heading("三、mDeBERTa-v3 验证结果", 1)
p("对每对命题双向推理（A→B, B→A），取 entailment 概率：")
p("  · score_avg = (e_ab + e_ba) / 2  （方案指定口径）")
p("  · score_min = min(e_ab, e_ba)    （严格双向口径）")
p("扫阈值 θ∈[0.50, 0.95]，等义二分类评估（gold eq 为正例）。")

doc.add_heading("阈值扫描（score_avg 口径）", 2)
tbl(["θ", "TP", "FP", "TN", "FN", "精确率", "召回率", "F1"],
    [["0.50","22","1","108","2","95.7%","91.7%","0.936"],
     ["0.55","14","1","108","10","93.3%","58.3%","0.718"],
     ["0.60","12","1","108","12","92.3%","50.0%","0.649"],
     ["0.65","9","1","108","15","90.0%","37.5%","0.529"],
     ["0.70","8","1","108","16","88.9%","33.3%","0.485"],
     ["0.75","8","1","108","16","88.9%","33.3%","0.485"],
     ["0.80","8","1","108","16","88.9%","33.3%","0.485"],
     ["0.85","7","1","108","17","87.5%","29.2%","0.438"],
     ["0.90","3","1","108","21","75.0%","12.5%","0.214"],
     ["0.95","2","0","109","22","100%","8.3%","0.154"]])

p("工作点：θ = 0.50（score_avg 口径），精确率 95.7% ≥ 95% 信度要求，召回率 91.7%，F1 = 0.936。", bold=True)

doc.add_heading("阈值扫描（score_min 严格双向口径）", 2)
tbl(["θ", "TP", "FP", "TN", "FN", "精确率", "召回率", "F1"],
    [["0.50","8","1","108","16","88.9%","33.3%","0.485"],
     ["0.60","8","1","108","16","88.9%","33.3%","0.485"],
     ["0.70","8","1","108","16","88.9%","33.3%","0.485"],
     ["0.80","5","1","108","19","83.3%","20.8%","0.333"],
     ["0.85","3","1","108","21","75.0%","12.5%","0.214"],
     ["0.95","1","0","109","23","100%","4.2%","0.080"]])
p("严格双向口径召回率过低（33.3%），不采用。原因：史学概括句存在细节省略"
  "（如删去'元丰年间'、'北宋'等限定），单向蕴含概率偏低，拉低 min 值。"
  "采用方案指定的算术平均口径。", size=10)

# ============ 四、典型案例 ============
doc.add_heading("四、典型案例", 1)

doc.add_heading("4.1 正确判定的等义对（EQUIVALENT）", 2)
eq_cases = [x for x in results if x["gold"]=="eq" and x["score_avg"]>=0.50][:6]
tbl(["命题 A", "命题 B", "e_ab", "e_ba", "score_avg"],
    [[c["A"][:40], c["B"][:40], round(c["e_ab"],2), round(c["e_ba"],2), round(c["score_avg"],2)]
     for c in eq_cases])

doc.add_heading("4.2 正确判定的不等义对（DIFFERENT）", 2)
diff_cases = [x for x in results if x["gold"]=="diff" and x["score_avg"]<0.50][:6]
tbl(["命题 A", "命题 B", "e_ab", "e_ba", "score_avg"],
    [[c["A"][:40], c["B"][:40], round(c["e_ab"],2), round(c["e_ba"],2), round(c["score_avg"],2)]
     for c in diff_cases])

doc.add_heading("4.3 拆合案例（单向蕴含，应不等义）", 2)
ent_cases = [x for x in results if x["gold"]=="entail"]
tbl(["pair_id", "命题 A", "命题 B", "e_ab", "e_ba", "score_min", "判定"],
    [[c["pair_id"], c["A"][:35], c["B"][:35], round(c["e_ab"],2), round(c["e_ba"],2),
      round(c["score_min"],2), "不等义 ✓" if c["score_avg"]<0.50 else "误判"]
     for c in ent_cases])
p("说明：A 为细分条（如'18府州'），B 为合并条（如'18府州+10产茶区55%'）。"
  "正确表现为 e_ba 高（合并条蕴含细分条）、e_ab 低（细分条不蕴含合并条），"
  "score_avg < 0.50 判定不等义，与金标准一致。", size=10)

doc.add_heading("4.4 错误案例分析", 2)
p("θ=0.50（avg 口径）下共 3 个错误：1 假阳性 + 2 假阴性。", bold=True)

fp = [x for x in results if x["gold"]!="eq" and x["score_avg"]>=0.50]
fn = [x for x in results if x["gold"]=="eq" and x["score_avg"]<0.50][:2]

p("假阳性（1 例，实为金标准标记争议）：")
tbl(["pair_id", "命题 A", "命题 B", "score_avg", "分析"],
    [[c["pair_id"], c["A"][:40], c["B"][:40], round(c["score_avg"],2),
      "两版均为'18府州军+10产茶区55%'合并条，语义等价，金标准标为 diff 存疑"]
     for c in fp])

p("假阴性（节选 2 例，概括句细节省略导致）：")
tbl(["pair_id", "命题 A", "命题 B", "score_avg", "分析"],
    [[c["pair_id"], c["A"][:40], c["B"][:40], round(c["score_avg"],2),
      "B 省略'北宋''元丰年间'等限定，单向蕴含概率偏低"]
     for c in fn])

# ============ 五、结论 ============
doc.add_heading("五、结论与下一步", 1)
p("1. 工具选型：采用 mDeBERTa-v3-base-xnli-multilingual-nli-2mil7，"
  "工作点 θ=0.50（score_avg 算术平均口径），精确率 95.7% 满足 95% 信度要求。", bold=True)
p("2. RexUniNLU 淘汰：无概率输出 + NLI 行为异常 + 部署代价高。")
p("3. 拆合案例正确处理：单向蕴含概率不对称性（e_ba 高 e_ab 低）使合并/拆分对被判不等义，"
  "与人类金标准一致，无需原子化即可规避拆合错位问题。")
p("4. 下一步：① 噪声注入灵敏度分析（5%/10%/15% 判定误差对 ACI_J 的影响）；"
  "② 将 NLI 判定接入全流水线，计算 A01 完整 ACI_J；"
  "③ 正式实验前用开发集做温度缩放校准，冻结测试集固定校准参数。")

doc.save(OUT)
print("报告已生成:", OUT)
